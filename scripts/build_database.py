#!/usr/bin/env python3
from __future__ import annotations

import csv
import gzip
import hashlib
import json
import os
import re
import sqlite3
import sys
import time
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
DATABASE = DATA_DIR / "transactions.sqlite3"
META = DATA_DIR / "build_meta.local.json"

LEGACY_HEADERS = [
    "成交日期", "城市", "区域", "商圈", "小区", "户型", "朝向", "楼层",
    "面积（m²）", "挂牌价（万元）", "成交价（万元）", "成交单价（元）", "成交周期", "网页链接",
]
REQUIRED_HEADERS = {
    "成交日期", "区域", "商圈", "小区", "户型", "面积（m²）", "成交价（万元）", "成交单价（元）",
}


def discover_source() -> Path:
    configured = os.environ.get("BEIJING_HOUSE_SOURCE", "").strip()
    if configured:
        path = Path(configured).expanduser().resolve()
        if not path.exists():
            raise FileNotFoundError(f"BEIJING_HOUSE_SOURCE 指向的文件不存在：{path}")
        return path

    enhanced = sorted(
        ROOT.glob("outputs/**/北京成交数据_2018-2026地理增强完整版.csv.gz"),
        key=lambda path: path.stat().st_mtime_ns,
        reverse=True,
    )
    if enhanced:
        return enhanced[0]

    for path in (
        ROOT / "北京成交数据.xlsx",
        ROOT / "北京成交数据(435008条_2018.04.04-2025.08.01).xlsx",
        ROOT / "北京123.xlsx",
    ):
        if path.exists():
            return path
    raise FileNotFoundError("未找到北京成交数据源；可通过 BEIJING_HOUSE_SOURCE 指定 .csv.gz 或 .xlsx 文件")


def source_signature(path: Path) -> dict:
    stat = path.stat()
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return {
        "filename": path.name,
        "format": "csv.gz" if path.name.endswith(".csv.gz") else path.suffix.lower().lstrip("."),
        "size": stat.st_size,
        "mtime_ns": stat.st_mtime_ns,
        "sha256": digest.hexdigest(),
    }


def cache_is_current(source_path: Path) -> bool:
    if not DATABASE.exists() or not META.exists():
        return False
    try:
        meta = json.loads(META.read_text(encoding="utf-8"))
        stat = source_path.stat()
        saved = meta.get("source", {})
        return (
            saved.get("filename") == source_path.name
            and saved.get("size") == stat.st_size
            and saved.get("mtime_ns") == stat.st_mtime_ns
        )
    except (OSError, ValueError):
        return False


def parse_cycle(value) -> int | None:
    if value is None:
        return None
    match = re.search(r"(\d+)", str(value))
    return int(match.group(1)) if match else None


def parse_rooms(value) -> str:
    text = str(value or "").strip()
    match = re.match(r"(\d+)室", text)
    if match:
        rooms = int(match.group(1))
        return "4室及以上" if rooms >= 4 else f"{rooms}室"
    if "开间" in text:
        return "开间"
    return "其他"


def clean_text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none", "nat"} else text


def clean_float(value) -> float | None:
    text = clean_text(value)
    if not text:
        return None
    try:
        return float(text)
    except (TypeError, ValueError):
        return None


def normalize_date(value) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = clean_text(value).replace(".", "-").replace("/", "-")
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10]).isoformat()
    except ValueError:
        return None


def iter_csv_rows(path: Path):
    with gzip.open(path, "rt", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = set(reader.fieldnames or [])
        missing = sorted(REQUIRED_HEADERS - headers)
        if missing:
            raise ValueError(f"增强 CSV 缺少字段：{missing}")
        yield from reader


def iter_excel_rows(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows))
        if headers != LEGACY_HEADERS:
            raise ValueError(f"Excel 字段与预期不一致：{headers}")
        for row in rows:
            yield dict(zip(headers, row))
    finally:
        workbook.close()


def iter_source_rows(path: Path):
    if path.name.endswith(".csv.gz"):
        return iter_csv_rows(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return iter_excel_rows(path)
    raise ValueError(f"不支持的数据格式：{path.name}")


def declared_complete_month(source_path: Path, max_date: str) -> str:
    configured = os.environ.get("BEIJING_HOUSE_COMPLETE_THROUGH", "").strip()
    if re.fullmatch(r"\d{4}-\d{2}", configured):
        return configured
    if source_path.name == "北京成交数据_2018-2026地理增强完整版.csv.gz" and max_date >= "2026-08-01":
        return "2026-08"
    parsed = date.fromisoformat(max_date)
    next_month = (parsed.replace(day=28) + timedelta(days=4)).replace(day=1)
    month_end = next_month - timedelta(days=1)
    if parsed == month_end:
        return parsed.strftime("%Y-%m")
    return (parsed.replace(day=1) - timedelta(days=1)).strftime("%Y-%m")


def build(force: bool = False) -> dict:
    source_path = discover_source()
    if not force and cache_is_current(source_path):
        return json.loads(META.read_text(encoding="utf-8"))

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    temp_db = DATABASE.with_suffix(".sqlite3.tmp")
    temp_db.unlink(missing_ok=True)

    started = time.time()
    source = source_signature(source_path)
    connection = sqlite3.connect(temp_db)
    connection.execute("PRAGMA journal_mode=OFF")
    connection.execute("PRAGMA synchronous=OFF")
    connection.execute("PRAGMA temp_store=MEMORY")
    connection.execute(
        """
        CREATE TABLE transactions (
            id INTEGER PRIMARY KEY,
            sale_date TEXT NOT NULL,
            sale_month TEXT NOT NULL,
            district TEXT NOT NULL,
            business_area TEXT NOT NULL,
            community TEXT NOT NULL,
            layout TEXT NOT NULL,
            rooms TEXT NOT NULL,
            orientation TEXT NOT NULL,
            floor TEXT NOT NULL,
            area REAL NOT NULL,
            listing_price REAL,
            sale_price REAL NOT NULL,
            unit_price REAL NOT NULL,
            cycle_days INTEGER,
            discount_rate REAL,
            url TEXT NOT NULL,
            source_name TEXT NOT NULL,
            source_record_id TEXT,
            merge_status TEXT,
            area_fill_method TEXT,
            business_fill_method TEXT,
            location_confidence TEXT
        )
        """
    )

    insert_sql = """
        INSERT INTO transactions (
            sale_date, sale_month, district, business_area, community, layout, rooms,
            orientation, floor, area, listing_price, sale_price, unit_price,
            cycle_days, discount_rate, url, source_name, source_record_id, merge_status,
            area_fill_method, business_fill_method, location_confidence
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """
    seen_keys: set[tuple[str, str]] = set()
    buffer = []
    stats = {
        "source_rows": 0,
        "kept_rows": 0,
        "excluded_duplicate": 0,
        "excluded_missing_location": 0,
        "excluded_unknown_district": 0,
        "excluded_outlier": 0,
        "excluded_bad_date": 0,
        "missing_listing_price": 0,
        "missing_source_url": 0,
    }
    source_counts: Counter[str] = Counter()
    merge_counts: Counter[str] = Counter()
    confidence_counts: Counter[str] = Counter()
    month_counts: Counter[str] = Counter()
    min_date = None
    max_date = None
    districts: set[str] = set()
    business_areas: set[str] = set()
    communities: set[str] = set()

    try:
        for row in iter_source_rows(source_path):
            stats["source_rows"] += 1
            sale_date = normalize_date(row.get("成交日期"))
            district = clean_text(row.get("区域"))
            business_area = clean_text(row.get("商圈"))
            community = clean_text(row.get("小区"))
            layout = clean_text(row.get("户型"))
            orientation = clean_text(row.get("朝向"))
            floor = clean_text(row.get("楼层"))
            url = clean_text(row.get("网页链接"))
            source_name = clean_text(row.get("原始来源")) or ("链家/贝壳" if url else "未知来源")
            source_record_id = clean_text(row.get("来源记录ID"))
            merge_status = clean_text(row.get("合并状态"))
            area_fill_method = clean_text(row.get("区域补全方法")) or "来源原值"
            business_fill_method = clean_text(row.get("商圈补全方法")) or "来源原值"
            location_confidence = clean_text(row.get("补全置信等级")) or "来源原值"

            if sale_date is None:
                stats["excluded_bad_date"] += 1
                continue
            if not district or not business_area or not community:
                stats["excluded_missing_location"] += 1
                continue
            if district == "-":
                stats["excluded_unknown_district"] += 1
                continue

            dedupe_key = ("url", url) if url else ((source_name, source_record_id) if source_record_id else None)
            if dedupe_key is not None:
                if dedupe_key in seen_keys:
                    stats["excluded_duplicate"] += 1
                    continue
                seen_keys.add(dedupe_key)

            area = clean_float(row.get("面积（m²）"))
            listing_price = clean_float(row.get("挂牌价（万元）"))
            sale_price = clean_float(row.get("成交价（万元）"))
            unit_price = clean_float(row.get("成交单价（元）"))
            if area is None or sale_price is None or unit_price is None:
                stats["excluded_outlier"] += 1
                continue
            if not (10 <= area <= 500 and 10 <= sale_price <= 10000 and 3000 <= unit_price <= 200000):
                stats["excluded_outlier"] += 1
                continue
            if listing_price is None or listing_price <= 0:
                stats["missing_listing_price"] += 1
                listing_price = None
                discount_rate = None
            else:
                discount_rate = sale_price / listing_price - 1
            if not url:
                stats["missing_source_url"] += 1

            month = sale_date[:7]
            min_date = sale_date if min_date is None or sale_date < min_date else min_date
            max_date = sale_date if max_date is None or sale_date > max_date else max_date
            districts.add(district)
            business_areas.add(f"{district}\t{business_area}")
            communities.add(f"{district}\t{business_area}\t{community}")
            source_counts[source_name] += 1
            merge_counts[merge_status or "未标注"] += 1
            confidence_counts[location_confidence] += 1
            month_counts[month] += 1
            buffer.append((
                sale_date, month, district, business_area, community, layout, parse_rooms(layout),
                orientation, floor, area, listing_price, sale_price, unit_price,
                parse_cycle(row.get("成交周期")), discount_rate, url, source_name, source_record_id or None,
                merge_status or None, area_fill_method, business_fill_method, location_confidence,
            ))
            stats["kept_rows"] += 1
            if len(buffer) >= 5000:
                connection.executemany(insert_sql, buffer)
                buffer.clear()

        if buffer:
            connection.executemany(insert_sql, buffer)
        connection.executescript(
            """
            CREATE INDEX idx_transactions_month ON transactions(sale_month);
            CREATE INDEX idx_transactions_district_month ON transactions(district, sale_month);
            CREATE INDEX idx_transactions_community_month ON transactions(community, sale_month);
            CREATE INDEX idx_transactions_business_month ON transactions(business_area, sale_month);
            CREATE INDEX idx_transactions_rooms_month ON transactions(rooms, sale_month);
            ANALYZE;
            """
        )
        connection.commit()
    finally:
        connection.close()

    if min_date is None or max_date is None:
        temp_db.unlink(missing_ok=True)
        raise ValueError("清洗后没有可用成交记录")

    if DATABASE.exists():
        DATABASE.unlink()
    temp_db.replace(DATABASE)
    meta = {
        "source": source,
        "built_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "elapsed_seconds": round(time.time() - started, 1),
        "date_min": min_date,
        "date_max": max_date,
        "default_end_month": declared_complete_month(source_path, max_date),
        "district_count": len(districts),
        "business_area_count": len(business_areas),
        "community_count": len(communities),
        "cleaning": stats,
        "source_counts": dict(source_counts.most_common()),
        "merge_status_counts": dict(merge_counts.most_common()),
        "location_confidence_counts": dict(confidence_counts.most_common()),
        "monthly_counts": dict(sorted(month_counts.items())),
        "geography": {
            "has_coordinates": False,
            "district_and_business_area_enhanced": source["format"] == "csv.gz",
            "note": "增强数据包含区域与商圈补全，但不包含经纬度；小区点位仍需通过高德 POI 搜索。",
        },
        "rules": {
            "area": [10, 500],
            "unit_price": [3000, 200000],
            "sale_price": [10, 10000],
            "location": "区域、商圈和小区均需非空",
        },
    }
    META.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    return meta


if __name__ == "__main__":
    result = build(force="--force" in sys.argv)
    print(json.dumps(result, ensure_ascii=False, indent=2))
